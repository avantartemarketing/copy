"""Build tool/release-copy.xlsx: one release's inputs -> fixed lines -> every post and email, all by formula."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import datetime, math

wb = Workbook()
ARIAL = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True); TITLE = Font(name="Arial", size=13, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF"); NOTE = Font(name="Arial", size=9, italic=True, color="555555")
YELLOW = PatternFill("solid", fgColor="FFFF00"); GREY = PatternFill("solid", fgColor="EDEDED")
WRAP = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
NL = "CHAR(10)"; PP = "CHAR(10)&CHAR(10)"

# ================================================================= Inputs
ws = wb.active; ws.title = "Inputs"
ws["A1"] = "One release · inputs"; ws["A1"].font = TITLE
ws["A2"] = "Blue cells on yellow are yours to fill. Everything else in this workbook is a formula. Set mechanic to draw to see the LE versions, and window to 48 hours or one week for the TL flow."; ws["A2"].font = NOTE
for c, h in zip("ABC", ["Field", "Value", "What it is"]): ws[f"{c}3"] = h; ws[f"{c}3"].font = BOLD; ws[f"{c}3"].fill = GREY
facts = [
    ("artist", "Grayson Perry", "full name"),
    ("handle", "alanmeasles", "Instagram handle, blank if none"),
    ("hashtag", "GraysonPerry", "announcement post only; blank for none"),
    ("ordinal", "latest", "debut, latest or second"),
    ("title", "The Changeling", "exactly as it appears; for a series with one shared title, give the series name"),
    ("is_series", "no", "yes when the works share one title"),
    ("edition_phrase", "a new limited edition print", "e.g. a trio of limited edition prints; six new limited edition prints"),
    ("works", 1, "number of works"),
    ("mechanic", "timed", "timed (TL) or draw (LE)"),
    ("window", "one week", "48 hours or one week; timed only"),
    ("launch_at", datetime.datetime(2026, 4, 30, 14, 0), "date and time the window opens, or the draw opens (UK)"),
    ("closes_at", datetime.datetime(2026, 5, 7, 14, 0), "date and time the window or draw closes (UK)"),
    ("beneficiary", "the Foundling Museum", "as it reads mid-sentence; blank if none"),
    ("beneficiary_handle", "foundlingmuseum", "blank if none"),
    ("advisor", "Sam", "signs the Insiders email and the first-time collector survey: Sofiya, Curtis or Sam"),
    ("early_access_code", "000-000", "draw releases"),
    ("feature_1", "Signed by the artist", "the standard three; clear a cell to drop it"),
    ("feature_2", "Individually numbered", ""),
    ("feature_3", "Free worldwide shipping", ""),
    ("feature_4", "", "bespoke, e.g. Hand-finished by the artist. Never 'our'."),
    ("feature_5", "", "bespoke, e.g. Available individually or as a set"),
]
R = {}; row = 4
for k, v, note in facts:
    ws[f"A{row}"] = k; ws[f"B{row}"] = v; ws[f"C{row}"] = note
    ws[f"A{row}"].font = ARIAL; ws[f"B{row}"].font = BLUE; ws[f"B{row}"].fill = YELLOW; ws[f"C{row}"].font = NOTE
    if isinstance(v, datetime.datetime): ws[f"B{row}"].number_format = "dd mmm yyyy hh:mm"
    R[k] = row; row += 1
row += 1; ws[f"A{row}"] = "Fragments, written once per release. Voice-neutral: no we or our, no artist name in the hook or making line."; ws[f"A{row}"].font = BOLD; row += 1
frags = [
    ("bio", "Grayson Perry has long been one of the most distinctive voices in contemporary British art.\n\nWorking across ceramics, tapestry, print and sculpture, his practice examines identity, class, and gender in modern society. Drawing on autobiography as well as social observation, Grayson often pairs the decorative language of traditional craft with direct commentary on the structures and contradictions that shape everyday life.", "2 to 3 sentences on the artist, third person, no handle. Coming Soon post only."),
    ("hook", "The Changeling portrays a child as a symbol of modern youth, exploring how technology shapes today's generations. It reimagines changeling folklore for the digital age, and is intended to be 'quite disturbing'.", "1 to 2 sentences about the work only. Reused in every post and email."),
    ("making", "To create the edition, printmakers at Make-Ready translated a new design into a 9-colour silkscreen, capturing the vivid detail of the original drawing.", "one sentence, written as 'printmakers at Make-Ready'; the sheet adds 'our'."),
    ("quote", "Beauty and seriousness are perhaps the most shocking tactics left to artists these days.", "verbatim, a complete sentence, or blank"),
]
for k, v, note in frags:
    ws[f"A{row}"] = k; ws[f"B{row}"] = v; ws[f"C{row}"] = note
    ws[f"A{row}"].font = ARIAL; ws[f"B{row}"].font = BLUE; ws[f"B{row}"].fill = YELLOW; ws[f"B{row}"].alignment = WRAP; ws[f"C{row}"].font = NOTE; ws[f"C{row}"].alignment = WRAP
    ws.row_dimensions[row].height = 16 * max(2, math.ceil(len(v) / 80) + v.count("\n")); R[k] = row; row += 1
row += 1; ws[f"A{row}"] = "Derived (formulas, do not edit)"; ws[f"A{row}"].font = BOLD; row += 1
B = lambda k: f"$B${R[k]}"
derived = ["named", "work", "Work", "work_email", "Work_email", "work_intro", "support", "beneficiary_tagged", "bio_tagged", "hook_names_beneficiary",
           "features", "features_list", "features_list_support", "making_ours", "making_ours_tagged",
           "launch_date", "launch_time", "launch_weekday", "launch_date_dd", "close_date", "close_time", "close_weekday", "close_date_dd",
           "collect_phrase", "edition_phrase_cap", "window_cap", "card_line", "card", "quote_line", "each_artwork", "work_word", "add_what"]
for k in derived:
    R[k] = row; ws[f"A{row}"] = k; ws[f"A{row}"].font = ARIAL; ws[f"B{row}"].font = ARIAL; ws[f"B{row}"].alignment = WRAP; row += 1
F = {
 "named": f'={B("artist")}&IF({B("handle")}="",""," (@"&{B("handle")}&")")',
 "work": f'=IF({B("is_series")}="yes","the "&{B("title")}&" series","‘"&{B("title")}&"’")',
 "Work": f'=UPPER(LEFT({B("work")},1))&MID({B("work")},2,999)',
 "work_email": f'=IF({B("is_series")}="yes","the "&{B("title")}&" series",{B("title")})',
 "Work_email": f'=UPPER(LEFT({B("work_email")},1))&MID({B("work_email")},2,999)',
 "work_intro": f'=IF({B("is_series")}="yes",{B("edition_phrase")}&" from the "&{B("title")}&" series",IF({B("works")}=1,{B("title")}&", "&{B("edition_phrase")},{B("edition_phrase")}))',
 "support": f'=IF(OR({B("beneficiary")}="",{B("hook_names_beneficiary")}),"",", released in support of "&{B("beneficiary")})',
 "beneficiary_tagged": f'=IF({B("beneficiary")}="","",{B("beneficiary")}&IF({B("beneficiary_handle")}="",""," (@"&{B("beneficiary_handle")}&")"))',
 "bio_tagged": f'=IF({B("handle")}="",{B("bio")},SUBSTITUTE({B("bio")},{B("artist")},{B("named")},1))',
 "each_artwork": f'=IF({B("works")}>1," for each artwork","")',
 "work_word": f'=IF({B("works")}>1,"works","work")',
 "add_what": f'=IF({B("works")}>1,"a print",{B("work_email")})',
 "hook_names_beneficiary": f'=IF({B("beneficiary")}="",FALSE,ISNUMBER(SEARCH({B("beneficiary")},{B("hook")})))',
 "features": f'={B("feature_1")}&IF({B("feature_2")}="","",". "&{B("feature_2")})&IF({B("feature_3")}="","",". "&{B("feature_3")})&IF({B("feature_4")}="","",". "&{B("feature_4")})&IF({B("feature_5")}="","",". "&{B("feature_5")})&"."',
 "features_list": f'="· "&{B("feature_1")}&IF({B("feature_2")}="","",{NL}&"· "&{B("feature_2")})&IF({B("feature_3")}="","",{NL}&"· "&{B("feature_3")})&IF({B("feature_4")}="","",{NL}&"· "&{B("feature_4")})&IF({B("feature_5")}="","",{NL}&"· "&{B("feature_5")})',
 "features_list_support": f'={B("features_list")}&IF(OR({B("beneficiary")}="",{B("hook_names_beneficiary")}),"",{NL}&"· Released in support of "&{B("beneficiary")})',
 "making_ours": f'=SUBSTITUTE({B("making")},"printmakers at Make-Ready","our printmakers at Make-Ready")',
 "making_ours_tagged": f'=SUBSTITUTE({B("making")},"printmakers at Make-Ready","our printmakers at Make-Ready (@make__ready)")',
 "launch_date": f'=TEXT({B("launch_at")},"d mmmm")', "launch_time": f'=TEXT({B("launch_at")},"hh:mm")&" UK time"', "launch_weekday": f'=TEXT({B("launch_at")},"dddd")', "launch_date_dd": f'=TEXT({B("launch_at")},"dd mmmm")',
 "close_date": f'=TEXT({B("closes_at")},"d mmmm")', "close_time": f'=TEXT({B("closes_at")},"hh:mm")&" UK time"', "close_weekday": f'=TEXT({B("closes_at")},"dddd")', "close_date_dd": f'=TEXT({B("closes_at")},"dd mmmm")',
 "collect_phrase": f'=SUBSTITUTE({B("edition_phrase")}," new "," ")',
 "edition_phrase_cap": f'=UPPER(LEFT({B("edition_phrase")},1))&MID({B("edition_phrase")},2,999)',
 "window_cap": f'=UPPER(LEFT({B("window")},1))&MID({B("window")},2,999)',
 "card_line": f'={B("edition_phrase_cap")}&" by "&{B("artist")}&"."',
 "card": f'={B("title")}&{NL}&{B("card_line")}',
 "quote_line": f'=IF({B("quote")}="","","“"&{B("quote")}&"” – "&{B("artist")})',
}
for k, f in F.items(): ws[f"B{R[k]}"] = f
ws.row_dimensions[R["bio_tagged"]].height = 80; ws.row_dimensions[R["features_list"]].height = 60; ws.row_dimensions[R["features_list_support"]].height = 70
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 78; ws.column_dimensions["C"].width = 62

# ================================================================= Lines
ls = wb.create_sheet("Lines")
ls["A1"] = "The fixed sentences, with the placeholders they take"; ls["A1"].font = TITLE
ls["A2"] = "This is the whole system's copy, for the posts and the emails. Edit a blue cell to change every future release. Column C fills the placeholders from Inputs."; ls["A2"].font = NOTE
for c, h in zip("ABC", ["Key", "Line, with placeholders", "Filled from Inputs"]): ls[f"{c}3"] = h; ls[f"{c}3"].font = BOLD; ls[f"{c}3"].fill = GREY
TOKEN = "{{ personalization_token('contact.firstname', 'there') }}"
lines = [
 ("POSTS", None),
 ("post · coming soon · status", "Our {ordinal} collaboration is on the horizon."),
 ("post · announce · status", "Announcing {work}, {edition_phrase} by {named}."),
 ("post · announce · status, series", "Announcing {edition_phrase} by {named}, from {work}."),
 ("post · sustain · status", "A closer look at {work}, {edition_phrase} by {named}."),
 ("post · artist's words · status", "“{quote}” – {artist}"),
 ("post · artist's words · second line", "{Work}, {edition_phrase} by {named}."),
 ("post · live · status", "{Work} by {named} is now available to collect, for {window} only."),
 ("post · still time · status", "There's still time to collect {work} by {named}."),
 ("post · last chance · status", "Last chance to collect {work} by {named}."),
 ("post · last chance · status, draw", "Last chance to enter the draw for {work} by {named}."),
 ("post · features line", "{features}"),
 ("post · beneficiary line", "Released in support of {beneficiary_tagged}."),
 ("post · action before launch, timed", "Available for {window} from {launch_time} on {launch_date}. Link in bio to get updates."),
 ("post · action before launch, draw", "Enter the draw via our link in bio. Closes {close_date} at {close_time}."),
 ("post · action while open, timed", "Link in bio to buy a print. Closes {close_date} at {close_time}."),
 ("post · action while open, draw", "Enter the draw via our link in bio. Closes {close_date} at {close_time}."),
 ("post · coming soon · action", "Link in bio to get updates."),
 ("post · hashtag", "#{hashtag}"),
 ("EMAILS · shared", None),
 ("email · greeting", "Hi " + TOKEN + ","),
 ("email · questions line", "If you have any questions, please don't hesitate to get in touch by replying to this email."),
 ("email · draw line", "Enter the draw for a chance to collect. Closes at {close_time} on {close_date}."),
 ("email · edition numbers line", "Earlier orders will typically receive a lower edition number, with framed editions receiving lower edition numbers than unframed prints. You can read more on how we allocate edition numbers across a release here."),
 ("EMAILS · announcement", None),
 ("email · announce · subject", "{artist} – {edition_phrase_cap}"),
 ("email · announce · opener", "We're delighted to announce our {ordinal} collaboration with {artist} – {work_intro}{support}."),
 ("email · announce · launch line, timed", "The edition will be available to collect for {window} only, starting at {launch_time} on {launch_weekday}, {launch_date_dd}."),
 ("email · announce · register line", "Click below to learn more and register for updates."),
 ("EMAILS · welcome (TL flow)", None),
 ("email · welcome · 1", "Welcome {{ personalization_token('contact.firstname', 'to Avant Arte') }}!"),
 ("email · welcome · 2", "Avant Arte began with a simple mission – to make collecting art more accessible. Since then, we've collaborated with hundreds of inspiring artists, from rising stars to icons like Ai Weiwei, Jenny Holzer, Lee Ufan and Carrie Mae Weems."),
 ("email · welcome · 3", "Our upcoming collaboration with {artist} is the latest in this lineage. If you're new to collecting art or curious about limited editions, our library of guides is a good place to start. In particular, How to collect art and What is an edition?"),
 ("email · welcome · 4", "Let us know if you have any questions. We're excited to see what you collect."),
 ("EMAILS · early access", None),
 ("email · early access, timed · opener", "Our {ordinal} collaboration with {artist} launches tomorrow at {launch_time} and will be available to collect for {window} only."),
 ("email · early access, timed · registered line", "As a thank you for registering for updates, we're offering you the chance to collect the release 24 hours before everyone else."),
 ("email · early access, timed · past collectors line", "As a previous collector of the artist, we're offering you the chance to order the collaboration 24 hours before everyone else."),
 ("email · early access, timed · unlock line", "Unlock early access using the link below."),
 ("email · early access, draw · 1", "We're currently preparing the launch of our {ordinal} collaboration with {artist} – {work_intro}{support}."),
 ("email · early access, draw · 2", "The edition will launch publicly on {launch_date} and will be allocated by a randomised draw; however, for the next 48 hours we're offering a small group of collectors a first look, plus access to a limited number of pre-orders. Based on your order history and the artists you've expressed an interest in, I thought the edition would be a good fit for your collection."),
 ("email · early access, draw · 3", "Explore the artwork and place your order via the private link below. Use code {early_access_code} to unlock early access."),
 ("email · insiders · opener", "I'm delighted to share our {ordinal} collaboration with {artist} – {work_intro}{support}."),
 ("email · insiders · mechanics, draw", "The edition will launch publicly on {launch_date} and will be allocated by a randomised draw. However, for the next 48 hours, I'm excited to offer you a first look, plus early access to a limited number of pre-orders."),
 ("email · insiders · mechanics, timed", "The edition launches tomorrow at {launch_time} and will be available to collect for {window} only. For the next 24 hours, I'm excited to offer you early access ahead of the public launch."),
 ("email · insiders · code line, draw", "Discover the artwork and place your order via the private link below. Use code {early_access_code} to unlock early access."),
 ("email · insiders · link line, timed", "Discover the artwork and place your order via the private link below."),
 ("email · insiders · sign-off", "Best regards,\n{advisor}\n\nArt Advisor at Avant Arte"),
 ("EMAILS · the window (TL)", None),
 ("email · live · opener", "Our {ordinal} collaboration with {artist} is now available to collect for {window} only – {work_intro}{support}."),
 ("email · live · closing line", "The opportunity to collect an edition ends at {close_time} on {close_weekday}, {close_date_dd}. Click the link below to add {add_what} to your collection."),
 ("email · halfway · kicker", "Collect a print by {artist}"),
 ("email · halfway · headline", "24 hours down, 24 to go"),
 ("email · halfway · footer", "There's still time to add a print to your collection."),
 ("email · 5 days · headline", "2 days down, 5 days to go"),
 ("email · 5 days · footer", "Add a signed print to your collection"),
 ("email · 3 days · headline", "Three days left to collect"),
 ("email · 3 days · 1", "There are just three days left to collect {collect_phrase} by {artist}."),
 ("email · 3 days · 2", "Click the link to collect before {close_time} on {close_weekday}, {close_date_dd}."),
 ("email · 3 days · footer", "Add a print to your collection"),
 ("EMAILS · last chance", None),
 ("email · last chance, timed · 1", "It's now or never for our {ordinal} collaboration with {artist} – {collect_phrase} will be available to order until {close_time} on {close_weekday}, {close_date_dd}."),
 ("email · last chance, timed · 2", "After this time, the edition size{each_artwork} will be confirmed, and the {work_word} will no longer be available to purchase."),
 ("email · last chance, draw · 1", "This is your final opportunity to enter the draw for {work_email}, {collect_phrase} by {artist}."),
 ("email · last chance, draw · 2", "For a chance to collect, click the link below to enter the draw. The draw closes at {close_time} on {close_weekday}, {close_date_dd}."),
 ("EMAILS · surveys (LE)", None),
 ("email · first-time survey · subject", "Congratulations on your first Avant Arte edition"),
 ("email · first-time survey · 1", "Thank you for adding our collaboration with {artist} to your collection. Great choice!"),
 ("email · first-time survey · 2", "I'm {advisor}, an art advisor at Avant Arte. I'm here to connect collectors with artists by making suggestions, answering questions and offering early access to our collaborations."),
 ("email · first-time survey · 3", "If you have two minutes to complete a short survey about your collecting journey to date, your responses will help guide my recommendations."),
 ("email · first-time survey · 4", "Any questions? Reply to this email."),
 ("email · non-purchaser survey · 1", "As someone who registered for updates but ultimately decided not to enter the draw for our collaboration with {artist}, we'd love your feedback."),
 ("email · non-purchaser survey · 2", "This two-minute survey will guide our future collaborations and help us recommend the right artists and editions for your collection."),
 ("email · non-purchaser survey · 3", "As a thank you for taking part in the survey and helping us improve the collector experience, you'll be entered into a draw to win a €500 Avant Arte gift card."),
 ("email · non-purchaser survey · terms", "You can read the terms here."),
 ("email · non-purchaser survey · footer", "What can we do better?"),
 ("EMAILS · monthly preview", None),
 ("email · monthly preview · opener", "From {artist} comes {work_intro}{support}."),
 ("email · monthly preview · when, draw", "The edition will be allocated by a randomised draw, which closes at {close_time} on {close_date}."),
 ("email · monthly preview · when, timed", "The edition will be available to collect for {window} only, from {launch_time} on {launch_weekday}, {launch_date_dd}."),
]
ph = {f"{{{k}}}": k for k in ["artist", "named", "work", "Work", "work_email", "Work_email", "work_intro", "support", "edition_phrase", "edition_phrase_cap", "collect_phrase", "ordinal", "window",
                             "launch_date", "launch_time", "launch_weekday", "launch_date_dd", "close_date", "close_time", "close_weekday", "close_date_dd",
                             "beneficiary_tagged", "quote", "hashtag", "features", "advisor", "early_access_code", "each_artwork", "work_word", "add_what"]}
L = {}; r = 4
for k, text in lines:
    if text is None:
        ls[f"A{r}"] = k; ls[f"A{r}"].font = BOLD; ls[f"A{r}"].fill = GREY; r += 1; continue
    L[k] = r; ls[f"A{r}"] = k; ls[f"B{r}"] = text
    ls[f"A{r}"].font = ARIAL; ls[f"B{r}"].font = BLUE; ls[f"B{r}"].fill = YELLOW; ls[f"B{r}"].alignment = WRAP; ls[f"C{r}"].alignment = WRAP; ls[f"C{r}"].font = ARIAL
    expr = f"B{r}"
    for p, field in ph.items():
        if p in text: expr = f'SUBSTITUTE({expr},"{p}",Inputs!$B${R[field]})'
    ls[f"C{r}"] = "=" + expr
    ls.row_dimensions[r].height = 16 * max(1, math.ceil(len(text) / 70) + text.count("\n"))
    r += 1
ls.column_dimensions["A"].width = 40; ls.column_dimensions["B"].width = 70; ls.column_dimensions["C"].width = 70

LC = lambda k: f"Lines!$C${L[k]}"
IB = lambda k: f"Inputs!$B${R[k]}"
draw = f'{IB("mechanic")}="draw"'
words_status = LC("post · artist's words · status"); words_second = LC("post · artist's words · second line")

# ================================================================= Posts
ps = wb.create_sheet("Posts")
ps["A1"] = "The posts, assembled"; ps["A1"].font = TITLE
ps["A2"] = "Every cell below is a formula over Inputs and Lines. Skeleton: status line · substance · features line · action line · hashtag. Coming Soon opens with the bio."; ps["A2"].font = NOTE
for i, h in enumerate(["#", "Post", "Made of", "Status line", "Substance", "Features line", "Action line", "Hashtag", "Caption", "Characters"], 1):
    c = ps.cell(row=3, column=i, value=h); c.font = BOLD; c.fill = GREY
feat = LC("post · features line")
benef = f'IF({IB("beneficiary_tagged")}="",""," "&{LC("post · beneficiary line")})'
pre = f'IF({draw},{LC("post · action before launch, draw")},{LC("post · action before launch, timed")})'
opn = f'IF({draw},{LC("post · action while open, draw")},{LC("post · action while open, timed")})'
def join(cells, guard=None, guard_text=None):
    expr = cells[0] + "".join(f'&IF({c}="","",{PP}&{c})' for c in cells[1:])
    return f'=IF({guard},"{guard_text}",{expr})' if guard else "=" + expr
quote = IB("quote")
NO_QUOTE = "(no quote given, post skipped)"
posts = [
 ("coming soon", "bio · status · action", f'={LC("post · coming soon · status")}', f'={IB("bio_tagged")}', '=""', f'={LC("post · coming soon · action")}', '=""', join(["E{r}", "D{r}", "G{r}"])),
 ("announce", "status · hook · features + beneficiary · action · hashtag", f'=IF({IB("is_series")}="yes",{LC("post · announce · status, series")},{LC("post · announce · status")})', f'={IB("hook")}', f'={feat}&IF({IB("hook_names_beneficiary")},"",{benef})', f'={pre}', f'=IF({IB("hashtag")}="","",{LC("post · hashtag")})', join(["D{r}", "E{r}", "F{r}", "G{r}", "H{r}"])),
 ("sustain", "status · making · features + beneficiary · action", f'={LC("post · sustain · status")}', f'={IB("making_ours_tagged")}', f'={feat}&{benef}', f'={pre}', '=""', join(["D{r}", "E{r}", "F{r}", "G{r}"])),
 ("sustain, the artist's words", "quote · status · features + beneficiary · action", f'=IF({quote}="","{NO_QUOTE}",{words_status})', f'=IF({quote}="","",{words_second})', f'=IF({quote}="","",{feat}&{benef})', f'=IF({quote}="","",{pre})', '=""', join(["D{r}", "E{r}", "F{r}", "G{r}"], f'{quote}=""', NO_QUOTE)),
 ("live (timed only)", "status · features + beneficiary · action", f'=IF({draw},"(not used for a draw)",{LC("post · live · status")})', '=""', f'=IF({draw},"",{feat}&{benef})', f'=IF({draw},"",{opn})', '=""', join(["D{r}", "F{r}", "G{r}"], draw, "(not used for a draw)")),
 ("still time (timed only)", "status · features + beneficiary · action", f'=IF({draw},"(not used for a draw)",{LC("post · still time · status")})', '=""', f'=IF({draw},"",{feat}&{benef})', f'=IF({draw},"",{opn})', '=""', join(["D{r}", "F{r}", "G{r}"], draw, "(not used for a draw)")),
 ("last chance", "status · features + beneficiary · action", f'=IF({draw},{LC("post · last chance · status, draw")},{LC("post · last chance · status")})', '=""', f'={feat}&{benef}', f'={opn}', '=""', join(["D{r}", "F{r}", "G{r}"])),
]
for i, (name, made, d, e, f, g, h, cap) in enumerate(posts):
    rr = 4 + i
    for col, v in enumerate([i + 1, name, made, d, e, f, g, h, cap.replace("{r}", str(rr)), f"=LEN(I{rr})"], 1):
        c = ps.cell(row=rr, column=col, value=v); c.font = ARIAL; c.alignment = WRAP; c.border = BOX
    ps.row_dimensions[rr].height = 190 if i in (0, 1) else 120
for col, w in zip("ABCDEFGHIJ", [4, 22, 34, 44, 56, 44, 44, 14, 92, 11]): ps.column_dimensions[col].width = w
ps.freeze_panes = "D4"

# ================================================================= Emails
es = wb.create_sheet("Emails")
es["A1"] = "The emails, assembled"; es["A1"].font = TITLE
es["A2"] = "Every cell below is a formula over Inputs and Lines. Body paragraphs are the fixed sentences with the same hook, making and quote dropped in. The 'Used?' column says when a row does not apply to this release."; es["A2"].font = NOTE
heads = ["#", "Email (comms plan name)", "For", "Used?", "Subject", "Preview", "Kicker", "Headline", "Body", "CTA", "Card", "Quote", "Footer", "Full email, ready to paste", "Characters"]
for i, h in enumerate(heads, 1):
    c = es.cell(row=3, column=i, value=h); c.font = BOLD; c.fill = GREY
A = IB("artist"); T = lambda s: f'"{s}"'
def P(*parts): return f'({parts[0]}' + "".join(f'&{PP}&{p}' for p in parts[1:]) + ")"
sel = lambda d_expr, t_expr: f'IF({draw},{d_expr},{t_expr})'
win48 = f'{IB("window")}="48 hours"'
USED = {
 "both": '"yes"',
 "timed": f'IF({draw},"(not used for a draw)","yes")',
 "timed 48h": f'IF({draw},"(not used for a draw)",IF({win48},"yes","(48-hour windows only; a one-week window gets the 5 days and 3 days emails)"))',
 "timed week": f'IF({draw},"(not used for a draw)",IF({win48},"(one-week windows only; a 48-hour window gets the halfway email)","yes"))',
 "draw": f'IF({draw},"yes","(not used for a timed edition)")',
}
E = ""  # empty cell
card = IB("card"); qline = IB("quote_line"); hook = IB("hook")
emails = [
 # name, for, used, subject, preview, kicker, headline, body, cta, card, quote, footer
 ("Announcement (TL Non-flow) · Announcement (LE)", "both", "both",
  LC("email · announce · subject"),
  sel(T("Enter the draw for a chance to collect. Closes ")+"&"+IB("close_date")+'&"."', T("Launching ")+"&"+IB("launch_date")+'&". Register for updates."'),
  sel(T("Enter the draw"), T("Launching ")+"&"+IB("launch_date")),
  sel(IB("Work_email"), IB("Work_email")+'&" by "&'+A),
  P(LC("email · announce · opener"), hook, IB("making_ours"), IB("features_list"), sel(LC("email · draw line"), P(LC("email · announce · launch line, timed"), LC("email · announce · register line")))),
  sel(T("Enter the draw"), T("Discover the collaboration")), card, qline,
  sel(T("Collect ")+"&"+IB("collect_phrase"), T("Launching ")+"&"+IB("launch_date"))),
 ("Welcome (TL Flow)", "timed", "timed", T("Welcome to Avant Arte"), T("Where the art world is more accessible."), E, E,
  P(LC("email · welcome · 1"), LC("email · welcome · 2"), LC("email · welcome · 3"), LC("email · welcome · 4")),
  T("Complete collector profile"), E, E, T("Where the art world is more accessible")),
 ("Early Access (TL Flow)", "timed", "timed", A+'&" – Early access 🔓"', T("Collect 24 hours before everyone else."), E, E,
  P(LC("email · greeting"), LC("email · early access, timed · opener"), hook, IB("features_list_support"), LC("email · early access, timed · registered line"), LC("email · early access, timed · unlock line")),
  T("Unlock early access"), E, E, T("24 hours ahead of the public launch")),
 ("Early access (TL Artist PP Non Flow)", "timed", "timed", A+'&" – Early access for past collectors"', T("Order before everyone else."), E, E,
  P(LC("email · greeting"), LC("email · early access, timed · opener"), hook, IB("features_list_support"), LC("email · early access, timed · past collectors line"), LC("email · early access, timed · unlock line"), LC("email · edition numbers line"), LC("email · questions line")),
  T("Unlock early access"), card, E, T("24 hours ahead of the public launch")),
 ("Early Access (TL Insiders) · Early Access (LE Insiders), signed by the advisor", "both", "both", A+'&" – Early access for Insiders 🔓"',
  sel(T("A first look, plus early access to a limited number of pre-orders."), T("Collect 24 hours before everyone else.")), E, E,
  P(LC("email · greeting"), LC("email · insiders · opener"), hook, sel(LC("email · insiders · mechanics, draw"), LC("email · insiders · mechanics, timed")), sel(LC("email · insiders · code line, draw"), LC("email · insiders · link line, timed")), LC("email · questions line"), LC("email · insiders · sign-off")),
  T("Unlock early access"), card, E, sel(T("Especially for you."), T("24 hours ahead of the public launch"))),
 ("Early/Exclusive access (LE)", "draw", "draw", A+'&" – Early access 🔓"', T("A first look, plus access to a limited number of pre-orders."), E, E,
  P(LC("email · greeting"), LC("email · early access, draw · 1"), hook, LC("email · early access, draw · 2"), LC("email · early access, draw · 3"), LC("email · edition numbers line"), LC("email · questions line")),
  T("Unlock early access"), card, E, T("Especially for you.")),
 ("Now Live (TL Flow) · Now Live (TL Non-flow)", "timed", "timed", A+'&" – Available now, for "&'+IB("window")+'&" only"', IB("window_cap")+'&", starting now."', T("Collect a print by ")+"&"+A, IB("window_cap")+'&", starting now"',
  P(LC("email · live · opener"), hook, IB("making_ours"), IB("features_list_support"), LC("email · live · closing line")),
  T("Buy now"), card, qline, T("Collect ")+"&"+IB("collect_phrase")),
 ("Halfway (TL Flow)", "timed, 48 hours", "timed 48h", A+'&" – 24 hours to go"', E, LC("email · halfway · kicker"), LC("email · halfway · headline"), E, T("Buy now"), E, E, LC("email · halfway · footer")),
 ("5 days to go (TL Flow)", "timed, one week", "timed week", A+'&" – 5 days left"', E, LC("email · halfway · kicker"), LC("email · 5 days · headline"), E, T("Buy a print"), E, E, LC("email · 5 days · footer")),
 ("3 days to go (TL Flow)", "timed, one week", "timed week", A+'&" – 3 days to go"', E, IB("Work_email")+'&" by "&'+A, LC("email · 3 days · headline"),
  P(LC("email · 3 days · 1"), hook, LC("email · 3 days · 2")), T("Buy a print"), card, E, LC("email · 3 days · footer")),
 ("Last chance (TL Flow) · Last chance (LE)", "both", "both", A+'&" – Last chance to "&'+sel(T("enter the draw"), T("collect")),
  sel(T("The draw closes at ")+"&"+IB("close_time")+'&" on "&'+IB("close_date")+'&"."', T("Time is almost up.")), A, sel(T("Last chance to enter the draw"), T("Last chance to collect")),
  sel(P(LC("email · last chance, draw · 1"), hook, LC("email · last chance, draw · 2")), P(LC("email · last chance, timed · 1"), LC("email · last chance, timed · 2"))),
  sel(T("Enter the draw"), T("Buy now")), card, E, T("Time is almost up")),
 ("First-time collector survey (LE)", "draw, after the draw closes", "draw", LC("email · first-time survey · subject"), E, E, E,
  P(LC("email · greeting"), LC("email · first-time survey · 1"), LC("email · first-time survey · 2"), LC("email · first-time survey · 3"), LC("email · first-time survey · 4")),
  T("Complete survey"), A, E, E),
 ("Non-purchaser survey (LE)", "draw, after the draw closes", "draw", A+'&" – Any feedback?"', E, T("Your feedback"), A,
  P(LC("email · non-purchaser survey · 1"), LC("email · non-purchaser survey · 2"), LC("email · non-purchaser survey · 3"), LC("email · non-purchaser survey · terms")),
  T("Complete survey"), E, E, LC("email · non-purchaser survey · footer")),
 ("Monthly Preview · this release's paragraph", "both; the advisor's monthly email is one paragraph per release, this is this release's", "both", E, E, E, E,
  P(LC("email · monthly preview · opener"), hook, IB("making_ours"), sel(LC("email · monthly preview · when, draw"), LC("email · monthly preview · when, timed"))),
  E, E, E, E),
]
for i, (name, use, used, subj, prev, kick, head, body, cta, cardf, quotef, foot) in enumerate(emails):
    rr = 4 + i
    guard = lambda expr: '=""' if expr == "" else f'=IF($D{rr}<>"yes","",{expr})'
    full = (f'=IF(D{rr}<>"yes",D{rr},IF(E{rr}="","","Subject: "&E{rr})&IF(F{rr}="","",{NL}&"Preview: "&F{rr})&IF(G{rr}="","",{NL}&"Kicker: "&G{rr})&IF(H{rr}="","",{NL}&"Headline: "&H{rr})'
            f'&IF(I{rr}="","",IF(E{rr}&F{rr}&G{rr}&H{rr}="","",{PP})&I{rr})&IF(J{rr}="","",{PP}&"CTA: "&J{rr})&IF(K{rr}="","",{PP}&"Card: "&K{rr})&IF(L{rr}="","",{PP}&"Quote: "&L{rr})&IF(M{rr}="","",{PP}&"Footer: "&M{rr}))')
    vals = [i + 1, name, use, "=" + USED[used]] + [guard(x) for x in (subj, prev, kick, head, body, cta, cardf, quotef, foot)] + [full, f"=LEN(N{rr})"]
    for col, v in enumerate(vals, 1):
        c = es.cell(row=rr, column=col, value=v); c.font = ARIAL; c.alignment = WRAP; c.border = BOX
    es.row_dimensions[rr].height = 300 if body else 90
for col, w in zip("ABCDEFGHIJKLMNO", [4, 30, 16, 16, 34, 30, 22, 26, 90, 16, 30, 34, 26, 96, 11]): es.column_dimensions[col].width = w
es.freeze_panes = "E4"

# ================================================================= Rules
rs = wb.create_sheet("Rules")
rs["A1"] = "The system on one page"; rs["A1"].font = TITLE
rules = [
 ("What it covers", "Every post on our own Instagram feed, and the emails in the comms plan: Announcement (TL and LE), Welcome, Early Access (TL Flow, TL Artist PP, Insiders TL and LE, LE early/exclusive), Now Live, Halfway, 5 days to go, 3 days to go, Last chance (TL and LE), the two LE surveys, and this release's paragraph for the Monthly Preview. The artist's own email is rare and stays hand-written."),
 ("Skeleton, posts", "Every post is: status line, substance, features line, action line, and a hashtag on the announcement only. Coming Soon opens with the bio."),
 ("Skeleton, emails", "Every email is: subject, preview, kicker, headline, body paragraphs, CTA, card, quote, footer. Bodies are fixed sentences from the Lines sheet with the same fragments dropped in. Halfway and 5 days to go have no body: they are image-led."),
 ("Fragments", "Four per release, written once: bio (2 to 3 sentences, Coming Soon only), hook (1 to 2 sentences about the work, in every post and email), making (one sentence naming Make-Ready), quote (verbatim). The hook and bio must not repeat each other."),
 ("Voice", "Fragments are voice-neutral: no we or our, no artist name in the hook or the making line, and 'printmakers at Make-Ready'. The sheet adds 'our'. The Insiders email and the first-time survey are signed by the advisor named on Inputs. Features never say 'our'."),
 ("Features", "The standard three (signed by the artist, individually numbered, free worldwide shipping) plus anything bespoke. One line of short sentences in a post, a list in an email. The beneficiary follows, named once."),
 ("Beneficiary", "One phrase everywhere: 'released in support of'. It goes in the opener unless the hook already names the beneficiary."),
 ("Action lines", "Three CTAs: get updates before launch, enter the draw for a draw, buy a print for an open window. One deadline form: Closes {date} at {time} UK time, or 'until {time} on {weekday}, {date}' in an email."),
 ("Mechanic and window", "Set mechanic to timed or draw and the sheet switches every line. For a timed edition, a 48-hour window gets the halfway email; a one-week window gets 5 days to go and 3 days to go."),
 ("Dates", "Enter launch and close once as date and time; the sheet derives '30 April', '14:00 UK time', 'Thursday' and '07 May' where each email needs them."),
 ("What stays written", "A delay reason, a framing paragraph, editorial deep dives, the rest of the Monthly Preview, the artist's own email, and anything for a second release."),
 ("How to use", "Fill the blue cells on Inputs. Read Posts and Emails. To change the house wording, edit the blue cells on Lines once."),
]
for i, (k, v) in enumerate(rules):
    rr = 3 + i; rs[f"A{rr}"] = k; rs[f"B{rr}"] = v; rs[f"A{rr}"].font = BOLD; rs[f"B{rr}"].font = ARIAL; rs[f"B{rr}"].alignment = WRAP; rs[f"A{rr}"].alignment = TOP
    rs.row_dimensions[rr].height = 46
rs.column_dimensions["A"].width = 20; rs.column_dimensions["B"].width = 112
wb.calculation.fullCalcOnLoad = True
import pathlib; wb.save(str(pathlib.Path(__file__).resolve().parent / "release-copy.xlsx")); print("saved", len(lines), "lines,", len(posts), "posts,", len(emails), "emails")
