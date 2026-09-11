"""Check that the workbook and the templates produce the same copy for a brief.

Evaluates every formula in the workbook with a small evaluator (the subset of Excel the sheet uses),
renders the templates for the same brief, and compares them line by line: posts, tweets and emails.

Usage:
  python3 tool/check_workbook.py                       # every brief in briefs/
  python3 tool/check_workbook.py briefs/joel-mesler-tl-26.yaml
"""
import datetime, pathlib, re, subprocess, sys

import yaml
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from render import render, email_phases, phases     # noqa: E402


# ----------------------------------------------------------------- a formula evaluator for the sheet's subset of Excel
class XLErr:
    def __bool__(self): raise ValueError("error used as condition")


TOK = re.compile(r'''\s*(?:
    (?P<str>"(?:[^"]|"")*")
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<ref>(?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!\$?[A-Z]{1,3}\$?\d+|\$?[A-Z]{1,3}\$?\d+)
  | (?P<name>[A-Za-z_][A-Za-z0-9_.]*)
  | (?P<op><>|>=|<=|=|>|<|&|\(|\)|,)
)''', re.X)


def tokenize(s):
    pos, out = 0, []
    while pos < len(s):
        m = TOK.match(s, pos)
        if not m or m.end() == pos:
            if s[pos:].strip() == "": break
            raise ValueError(f"bad token at {pos}: {s[pos:pos+30]!r}")
        pos = m.end(); out.append((m.lastgroup, m.group(m.lastgroup)))
    return out


class Evaluator:
    def __init__(self, wb):
        self.wb, self.cache = wb, {}
    def get(self, sheet, ref):
        ref = ref.replace("$", ""); key = (sheet, ref)
        if key not in self.cache:
            v = self.wb[sheet][ref].value
            self.cache[key] = self.eval(v[1:], sheet) if isinstance(v, str) and v.startswith("=") else ("" if v is None else v)
        return self.cache[key]
    def eval(self, formula, sheet):
        saved = getattr(self, "toks", None), getattr(self, "i", None), getattr(self, "sheet", None)
        self.toks, self.i, self.sheet = tokenize(formula), 0, sheet
        v = self.expr()
        if self.i != len(self.toks): raise ValueError(f"trailing tokens in {formula[:60]}")
        self.toks, self.i, self.sheet = saved
        return v
    def peek(self): return self.toks[self.i] if self.i < len(self.toks) else (None, None)
    def take(self, val=None):
        t = self.toks[self.i]; self.i += 1
        if val is not None and t[1] != val: raise ValueError(f"expected {val} got {t}")
        return t
    def expr(self):
        left = self.concat()
        while self.peek()[1] in ("=", "<>", ">", "<", ">=", "<="):
            op = self.take()[1]; right = self.concat()
            if op in ("=", "<>"):
                eq = self.equal(left, right); left = eq if op == "=" else not eq
            else:
                a, b = (left or 0), (right or 0); left = {">": a > b, "<": a < b, ">=": a >= b, "<=": a <= b}[op]
        return left
    @staticmethod
    def equal(a, b):
        a = "" if a is None else a; b = "" if b is None else b
        if isinstance(a, str) and isinstance(b, str): return a.lower() == b.lower()
        if isinstance(a, str) or isinstance(b, str): return (a == "" and b in (0, "")) or (b == "" and a in (0, ""))
        return a == b
    def concat(self):
        left = self.primary()
        while self.peek()[1] == "&":
            self.take(); left = self.text(left) + self.text(self.primary())
        return left
    @staticmethod
    def text(v):
        if v is None: return ""
        if isinstance(v, bool): return "TRUE" if v else "FALSE"
        if isinstance(v, (int, float)): return str(int(v)) if float(v).is_integer() else str(v)
        return str(v)
    def primary(self):
        kind, val = self.take()
        if kind == "str": return val[1:-1].replace('""', '"')
        if kind == "num": return float(val) if "." in val else int(val)
        if kind == "ref":
            sh, r = (val.split("!")[0].strip("'"), val.split("!")[1]) if "!" in val else (self.sheet, val)
            return self.get(sh, r)
        if kind == "name":
            name = val.upper()
            if name in ("TRUE", "FALSE"): return name == "TRUE"
            self.take("("); args = []
            if self.peek()[1] != ")":
                args.append(self.expr())
                while self.peek()[1] == ",": self.take(); args.append(self.expr())
            self.take(")")
            return self.call(name, args)
        if val == "(":
            v = self.expr(); self.take(")"); return v
        raise ValueError(f"unexpected {kind} {val}")
    def call(self, name, a):
        s = self.text
        if name == "IF": return a[1] if a[0] else (a[2] if len(a) > 2 else False)
        if name == "OR": return any(bool(x) for x in a)
        if name == "AND": return all(bool(x) for x in a)
        if name == "NOT": return not a[0]
        if name == "ISNUMBER": return isinstance(a[0], (int, float)) and not isinstance(a[0], bool)
        if name == "SEARCH": i = s(a[1]).lower().find(s(a[0]).lower()); return XLErr() if i < 0 else i + 1
        if name == "FIND": i = s(a[1]).find(s(a[0])); return XLErr() if i < 0 else i + 1
        if name == "LEFT": return s(a[0])[:int(a[1]) if len(a) > 1 else 1]
        if name == "MID": return s(a[0])[int(a[1]) - 1:int(a[1]) - 1 + int(a[2])]
        if name == "UPPER": return s(a[0]).upper()
        if name == "LEN": return len(s(a[0]))
        if name == "CHAR": return chr(int(a[0]))
        if name == "SUBSTITUTE":
            t, old, new = s(a[0]), s(a[1]), s(a[2])
            if old == "": return t
            if len(a) > 3:
                n = int(a[3]); parts = t.split(old)
                return t if len(parts) <= n else old.join(parts[:n]) + new + old.join(parts[n:])
            return t.replace(old, new)
        if name == "TEXT":
            d, fmt = a[0], s(a[1])
            codes = {"mmmm": d.strftime("%B"), "mmm": d.strftime("%b"), "dddd": d.strftime("%A"), "yyyy": d.strftime("%Y"), "hh": d.strftime("%H"),
                     "mm": d.strftime("%M") if "hh" in fmt else d.strftime("%m"), "dd": d.strftime("%d"), "d": str(d.day)}
            return re.sub(r"mmmm|mmm|dddd|yyyy|hh|mm|dd|d", lambda m: codes[m.group(0)], fmt)
        raise ValueError(f"unknown function {name}")


# ----------------------------------------------------------------- normalise both sides to plain lines
MARK = {"[SUBJECT]": "Subject:", "[PREVIEW]": "Preview:", "[KICKER]": "Kicker:", "[HEADLINE]": "Headline:", "[CTA]": "CTA:", "[CARD]": "Card:", "[FOOTER]": "Footer:"}


def lines_of(text):
    return [l.strip() for l in text.replace("\r", "").split("\n") if l.strip()]


def template_email_lines(chunk):
    out, quote = [], None
    for l in lines_of(chunk):
        if l.startswith(("[EMAIL", "[BODY]", "[FEATURES]", "[PARAGRAPH]")): continue
        if l.startswith("[QUOTE ATTRIBUTION]"): out.append(f"Quote: “{quote}” – {l.split('] ', 1)[1]}"); continue
        if l.startswith("[QUOTE]"): quote = l.split("] ", 1)[1]; continue
        for k, v in MARK.items():
            if l.startswith(k): l = v + l[len(k):]; break
        out.append(l)
    return out


def excel_email_lines(full):
    return lines_of(full)


def compare(name, a, b, report):
    if a == b:
        report.append(f"  ok    {name}")
        return True
    report.append(f"  DIFF  {name}")
    for i, (x, y) in enumerate(zip(a, b)):
        if x != y:
            report.append(f"          template: {x[:110]}"); report.append(f"          workbook: {y[:110]}"); break
    else:
        report.append(f"          line counts differ: template {len(a)}, workbook {len(b)}")
    return False


EMAIL_ROW = {"announce": 4, "welcome": 5, "early_access_tl": 6, "early_access_pp": 7, "insiders_ea": 8, "early_access_le": 9, "live": 10,
             "halfway": 11, "five_days": 12, "three_days": 13, "last_chance": 14, "survey_first": 15, "survey_nonpurchaser": 16, "monthly_preview": 17}
POST_ROW = {"coming soon": 4, "announce": 5, "sustain": 6, "sustain, the artist's words": 7, "live": 8, "still time": 9, "last chance": 10}
TWEET_ROW = {"coming soon": 4, "announce": 5, "live": 6, "still time": 7, "last chance": 8}


def check(brief_path):
    brief = yaml.safe_load(open(brief_path, encoding="utf-8"))
    stem = brief_path.stem
    xlsx = ROOT / "tool" / ("release-copy.xlsx" if stem == "grayson-perry-tl-26" else f"release-copy-{stem}.xlsx")
    subprocess.run([sys.executable, str(ROOT / "tool" / "build_release_copy.py"), str(brief_path), str(xlsx)], check=True, capture_output=True)
    ev = Evaluator(load_workbook(xlsx))
    report, ok = [f"{stem}  ({xlsx.name})"], True
    for ws in ev.wb.worksheets:                                       # every formula must evaluate
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = ev.get(ws.title, c.coordinate)
                    if isinstance(v, XLErr): report.append(f"  ERROR {ws.title}!{c.coordinate} evaluates to #VALUE!"); ok = False
    # posts
    posts = render(brief, "ig-set").split("\n[POST ")[0:]
    chunks = re.split(r"^\[POST \d+ · ", render(brief, "ig-set"), flags=re.M)[1:]
    for chunk in chunks:
        phase = chunk.split(" · ")[0]
        body = chunk.split("\n", 1)[1]
        ok &= compare(f"post · {phase}", lines_of(body), lines_of(ev.get("Posts", f"J{POST_ROW[phase]}")), report)
    chunks = re.split(r"^\[TWEET \d+ · ", render(brief, "twitter-set"), flags=re.M)[1:]
    for chunk in chunks:
        phase = chunk.split(" · ")[0]
        body = chunk.split("\n", 1)[1]
        ok &= compare(f"tweet · {phase}", lines_of(body), lines_of(ev.get("Tweets", f"H{TWEET_ROW[phase]}")), report)
    chunks = [c for c in render(brief, "email-set").split("════════════════════════════════════════════════════════════════") if c.strip()]
    keys = [k for k, _, _ in email_phases(brief["release"], brief["artist"])]
    for key, chunk in zip(keys, chunks):
        ok &= compare(f"email · {key}", template_email_lines(chunk), excel_email_lines(ev.get("Emails", f"N{EMAIL_ROW[key]}")), report)
    print("\n".join(report))
    return ok


if __name__ == "__main__":
    paths = [pathlib.Path(p) for p in sys.argv[1:]] or sorted((ROOT / "briefs").glob("*.yaml"))
    all_ok = all([check(p) for p in paths])
    print("ALL MATCH" if all_ok else "DIFFERENCES FOUND")
    sys.exit(0 if all_ok else 1)
